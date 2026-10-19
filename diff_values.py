import openpyxl

def norm(s):
    return "" if s is None else str(s).strip()

def is_num(v):
    if v is None: return False
    if isinstance(v,(int,float)): return True
    try: float(str(v).replace(",",".")); return True
    except: return False

def num(v):
    return float(str(v).replace(",",".")) if is_num(v) else None

MONTHS = ["JANEIRO","FEVEREIRO","MARÇO","ABRIL","MAIO","JUNHO","JULHO",
          "AGOSTO","SETEMBRO","OUTUBRO","NOVEMBRO","DEZEMBRO"]

def parse_sheet(ws):
    # header row
    hr=None
    for r in range(1,12):
        for c in range(1,ws.max_column+1):
            if norm(ws.cell(r,c).value).upper().startswith("PROCEDIMENTO"):
                hr=r;break
        if hr:break
    monthcol={}; tabcol={}; totalcol=None
    for c in range(1,ws.max_column+1):
        h=norm(ws.cell(hr,c).value).upper()
        if h in MONTHS: monthcol[h]=c
        elif "/TAB" in h:
            m=h.split("/")[0]
            # map abbrev to full
            for full in MONTHS:
                if full.startswith(m) or m.startswith(full[:3]):
                    tabcol[full]=c;break
        elif h=="TOTAL": totalcol=c
    # proc rows until 'total procedimentos'
    rows={}
    order=[]
    for r in range(hr+1, ws.max_row+1):
        lbl=norm(ws.cell(r,1).value)
        if lbl=="":continue
        if lbl.upper().startswith("TOTAL PROCEDIMENTOS") or lbl.upper().startswith("TOTAL DE PROCEDIMENTOS"):
            break
        # key by procedure code prefix (e.g. '208')
        code = lbl.split("-")[0].strip()
        rows[code]={"label":lbl,"row":r}
        order.append(code)
    return {"hr":hr,"monthcol":monthcol,"tabcol":tabcol,"rows":rows,"order":order,"ws":ws}

ANT=openpyxl.load_workbook("data/producao2026_v2.xlsx",data_only=True)
NEW=openpyxl.load_workbook("data/atualizada.xlsx",data_only=True)

for name in NEW.sheetnames:
    print("\n"+"#"*90)
    print("ABA:",name)
    print("#"*90)
    a=parse_sheet(ANT[name]); n=parse_sheet(NEW[name])
    # compare procedure codes present
    acodes=set(a["rows"]); ncodes=set(n["rows"])
    only_a=acodes-ncodes; only_n=ncodes-acodes
    if only_a: print("  [!] códigos SÓ na antiga:",sorted(only_a))
    if only_n: print("  [!] códigos SÓ na atualizada:",sorted(only_n))
    print("  [info] meses TAB disponíveis na antiga:", sorted(a["tabcol"].keys(), key=lambda m:MONTHS.index(m)))
    # compare realizado values month by month for shared codes
    diffs=[]
    for code in n["order"]:
        if code not in a["rows"]: continue
        ar=a["rows"][code]["row"]; nr=n["rows"][code]["row"]
        for m in MONTHS:
            ac=a["monthcol"].get(m); nc=n["monthcol"].get(m)
            if ac is None or nc is None: continue
            av=num(a["ws"].cell(ar,ac).value); nv=num(n["ws"].cell(nr,nc).value)
            if av is None and nv is None: continue
            if (av or 0)!=(nv or 0):
                diffs.append((code,m,av,nv,n["rows"][code]["label"][:35]))
    if diffs:
        print(f"  [DIFF] {len(diffs)} diferenças de REALIZADO (antiga -> atualizada):")
        for code,m,av,nv,lbl in diffs:
            print(f"       {code} {m:<10} {str(av):>7} -> {str(nv):>7}   ({lbl})")
    else:
        print("  [ok] REALIZADO idêntico entre as duas versões")

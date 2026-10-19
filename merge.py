import openpyxl
from copy import copy

def norm(s):
    return "" if s is None else str(s).strip()

def is_num(v):
    if v is None: return False
    if isinstance(v,(int,float)): return True
    try: float(str(v).replace(",",".")); return True
    except: return False

def num(v):
    return float(str(v).replace(",",".")) if is_num(v) else None

MONTHS = ["JANEIRO","FEVEREIRO","MARÇO","ABRIL","MAIO","JUNHO","JULHO",
          "AGOSTO","SETEMBRO","OUTUBRO","NOVEMBRO","DEZEMBRO"]
TAB_MONTHS = ["JANEIRO","FEVEREIRO","MARÇO","ABRIL","MAIO","JUNHO","JULHO"]

def find_header_row(ws):
    for r in range(1,12):
        for c in range(1,ws.max_column+1):
            if norm(ws.cell(r,c).value).upper().startswith("PROCEDIMENTO"):
                return r
    return None

def map_months(ws, hr):
    monthcol={}; tabcol={}; totalcol=None
    for c in range(1, ws.max_column+1):
        h = norm(ws.cell(hr,c).value).upper()
        if h in MONTHS:
            monthcol[h]=c
        elif "/TAB" in h:
            abbr = h.split("/")[0]
            for full in MONTHS:
                if full==abbr or full.startswith(abbr) or abbr.startswith(full[:3]):
                    tabcol[full]=c; break
        elif h=="TOTAL":
            totalcol=c
    return monthcol, tabcol, totalcol

def proc_rows(ws, hr):
    rows={}
    for r in range(hr+1, ws.max_row+1):
        lbl=norm(ws.cell(r,1).value)
        if lbl=="":continue
        if lbl.upper().startswith("TOTAL PROCEDIMENTOS") or lbl.upper().startswith("TOTAL DE PROCEDIMENTOS"):
            break
        code=lbl.split("-")[0].strip()
        rows[code]=r
    return rows

# ANTIGA: data_only for TAB numeric values
ANT = openpyxl.load_workbook("data/producao2026_v2.xlsx", data_only=True)
# BASE: editable (keeps styles) — we will FLATTEN formula cells to their computed values
BASE = openpyxl.load_workbook("data/atualizada.xlsx")
# CALC: computed values of the updated file (resolves formulas like =80+32 -> 112)
CALC = openpyxl.load_workbook("data/atualizada.xlsx", data_only=True)

ALL_SHEETS = BASE.sheetnames
CAPS_SHEETS = ["CAPSIJ","CAPS II","CAPS AD III"]

# STEP 1: flatten every formula cell in BASE to its computed value (from CALC),
# across all sheets, so no '=...' strings remain and all numbers are real.
flattened=0
formula_cells_without_cache=[]
for name in ALL_SHEETS:
    ws=BASE[name]; cws=CALC[name]
    for row in ws.iter_rows():
        for cell in row:
            v=cell.value
            if isinstance(v,str) and v.startswith("="):
                computed=cws.cell(cell.row, cell.column).value
                if computed is None:
                    formula_cells_without_cache.append((name,cell.coordinate,v))
                cell.value=computed
                flattened+=1
print(f"[flatten] {flattened} células de fórmula convertidas para valor fixo")
if formula_cells_without_cache:
    print("  [!] fórmulas SEM valor calculado em cache (viraram vazio):")
    for n,coord,f in formula_cells_without_cache:
        print(f"      {n}!{coord} = {f}")

# STEP 2: insert standardized TAB columns (jan..jul) from ANTIGA into CAPS sheets
report=[]
for name in CAPS_SHEETS:
    ws = BASE[name]
    hr = find_header_row(ws)
    base_month, _, _ = map_months(ws, hr)
    aws = ANT[name]; ahr = find_header_row(aws)
    _, a_tab, _ = map_months(aws, ahr)
    base_rows = proc_rows(ws, hr)
    a_rows = proc_rows(aws, ahr)

    # right-to-left insertion keeps earlier indices valid
    plan = sorted([(base_month[m], m) for m in MONTHS if m in base_month], reverse=True)
    inserted=0
    for col, m in plan:
        if m not in TAB_MONTHS:
            continue
        insert_at = col + 1
        ws.insert_cols(insert_at)
        hcell = ws.cell(hr, insert_at)
        hcell.value = f"{m}/TAB"
        src = ws.cell(hr, col)
        hcell.font=copy(src.font); hcell.fill=copy(src.fill)
        hcell.border=copy(src.border); hcell.alignment=copy(src.alignment)
        hcell.number_format=src.number_format
        for code, brow in base_rows.items():
            if code in a_rows and m in a_tab:
                av = aws.cell(a_rows[code], a_tab[m]).value
                if is_num(av):
                    ws.cell(brow, insert_at).value = float(av)
        inserted+=1
    report.append((name, inserted))

OUT="data/producao_final_merged.xlsx"
BASE.save(OUT)
print("Salvo:", OUT)
for name,ins in report:
    print(f"  [{name}] {ins} colunas TAB inseridas (jan..jul)")

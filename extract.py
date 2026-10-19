import openpyxl, json, re

def norm(s): return "" if s is None else str(s).strip()
def is_num(v):
    if v is None: return False
    if isinstance(v,(int,float)): return True
    try: float(str(v).replace(",",".")); return True
    except: return False
def num(v): return float(str(v).replace(",",".")) if is_num(v) else None

MONTHS = ["JANEIRO","FEVEREIRO","MARÇO","ABRIL","MAIO","JUNHO","JULHO",
          "AGOSTO","SETEMBRO","OUTUBRO","NOVEMBRO","DEZEMBRO"]

UNIT_NAMES = {
    "CAPSIJ": "CAPS Infanto-Juvenil",
    "CAPS II": "CAPS II",
    "CAPS AD III": "CAPS AD III",
    "UAA": "Unidade de Acolhimento Adulto",
    "UAI": "Unidade de Acolhimento Infantojuvenil",
}

def find_header_row(ws):
    for r in range(1,15):
        for c in range(1,ws.max_column+1):
            if norm(ws.cell(r,c).value).upper().startswith("PROCEDIMENTO"):
                return r
    return None

def map_cols(ws,hr):
    monthcol={}; tabcol={}
    for c in range(1,ws.max_column+1):
        h=norm(ws.cell(hr,c).value).upper()
        if h in MONTHS: monthcol[h]=c
        elif "/TAB" in h:
            a=h.split("/")[0]
            for f in MONTHS:
                if f==a or f.startswith(a) or a.startswith(f[:3]): tabcol[f]=c;break
    return monthcol,tabcol

def extract(path):
    wb=openpyxl.load_workbook(path,data_only=True)
    out={"units":[], "warnings":[]}
    for name in wb.sheetnames:
        ws=wb[name]
        hr=find_header_row(ws)
        monthcol,tabcol=map_cols(ws,hr)
        months_present=[m for m in MONTHS if m in monthcol and
                        any(is_num(ws.cell(r,monthcol[m]).value) for r in range(hr+1,ws.max_row+1))]
        procs=[]
        for r in range(hr+1, ws.max_row+1):
            lbl=norm(ws.cell(r,1).value)
            if lbl=="": continue
            up=lbl.upper()
            # STOP at the totals line -> everything below is ignored (per user instruction)
            if up.startswith("TOTAL PROCEDIMENTOS") or up.startswith("TOTAL DE PROCEDIMENTOS"):
                break
            # must start with an identification code (digits, optional leading zeros)
            m=re.match(r"^(\d{2,4})\s*-?\s*(.*)$", lbl)
            if not m:
                continue  # not a procedure row -> skip
            code=m.group(1)
            desc=m.group(2).strip()
            realizado={}; tab={}
            for mo in MONTHS:
                if mo in monthcol:
                    v=num(ws.cell(r,monthcol[mo]).value)
                    if v is not None: realizado[mo]=v
                if mo in tabcol:
                    tv=ws.cell(r,tabcol[mo]).value
                    if is_num(tv): tab[mo]=num(tv)
                    elif tv is not None and norm(tv)!="":
                        out["warnings"].append(f"{name} / {code} / {mo}/TAB: valor não numérico '{tv}'")
            # non-numeric realized check
            for mo in MONTHS:
                if mo in monthcol:
                    raw=ws.cell(r,monthcol[mo]).value
                    if raw is not None and not is_num(raw):
                        out["warnings"].append(f"{name} / {code} / {mo}: valor não numérico '{raw}' (tratado como ausente)")
            procs.append({"code":code,"desc":desc,"realizado":realizado,"tab":tab})
        has_tab=any(p["tab"] for p in procs)
        out["units"].append({
            "sheet":name,
            "name":UNIT_NAMES.get(name,name),
            "has_tab":has_tab,
            "months_present":months_present,
            "procedures":procs,
        })
    return out

if __name__=="__main__":
    import sys
    path=sys.argv[1] if len(sys.argv)>1 else "data/producao_final_merged.xlsx"
    data=extract(path)
    # summary
    print(f"Arquivo: {path}\n")
    for u in data["units"]:
        tot_real=sum(sum(p["realizado"].values()) for p in u["procedures"])
        tot_tab=sum(sum(p["tab"].values()) for p in u["procedures"])
        print(f"### {u['name']}  (aba {u['sheet']})")
        print(f"    procedimentos: {len(u['procedures'])} | tem TAB: {u['has_tab']}")
        print(f"    meses com dados: {u['months_present']}")
        print(f"    TOTAL realizado (recalculado): {tot_real:.0f}")
        if u['has_tab']:
            print(f"    TOTAL tabulado (recalculado): {tot_tab:.0f}  | taxa geral: {100*tot_tab/tot_real:.1f}%")
        print()
    if data["warnings"]:
        print("AVISOS DE VALIDAÇÃO:")
        for w in data["warnings"]:
            print("  -",w)
    else:
        print("AVISOS DE VALIDAÇÃO: nenhum")
    with open("data/dados_extraidos.json","w",encoding="utf-8") as f:
        json.dump(data,f,ensure_ascii=False,indent=2)
    print("\nJSON salvo em data/dados_extraidos.json")

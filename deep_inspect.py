import openpyxl, zipfile, os, json

PATH = "data/producao2026.xlsx"
wb = openpyxl.load_workbook(PATH, data_only=True)

print("="*90)
print("DUMP COMPLETO DE TODAS AS ABAS (todas as linhas, colunas com conteudo)")
print("="*90)

for name in wb.sheetnames:
    ws = wb[name]
    print(f"\n\n############ ABA: {name} (max_row={ws.max_row}, max_col={ws.max_column}) ############")
    for r in range(1, ws.max_row + 1):
        rowvals = []
        has_content = False
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=r, column=c).value
            if v is None:
                rowvals.append("")
            else:
                has_content = True
                rowvals.append(str(v))
        if has_content:
            # trim trailing empties
            while rowvals and rowvals[-1] == "":
                rowvals.pop()
            print(f"R{r:>2}| " + " | ".join(rowvals))

# Extract embedded images
print("\n\n" + "="*90)
print("IMAGENS EMBUTIDAS")
print("="*90)
os.makedirs("assets", exist_ok=True)
with zipfile.ZipFile(PATH) as z:
    for n in z.namelist():
        if n.startswith("xl/media/"):
            data = z.read(n)
            out = os.path.join("assets", os.path.basename(n))
            with open(out, "wb") as f:
                f.write(data)
            print(f"  {n} -> {out} ({len(data)} bytes)")

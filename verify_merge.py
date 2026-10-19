import openpyxl

def norm(s): return "" if s is None else str(s).strip()
def is_num(v):
    if v is None: return False
    if isinstance(v,(int,float)): return True
    try: float(str(v).replace(",",".")); return True
    except: return False
def num(v): return float(str(v).replace(",",".")) if is_num(v) else None

MONTHS=["JANEIRO","FEVEREIRO","MARÇO","ABRIL","MAIO","JUNHO","JULHO",
        "AGOSTO","SETEMBRO","OUTUBRO","NOVEMBRO","DEZEMBRO"]
TAB_MONTHS=["JANEIRO","FEVEREIRO","MARÇO","ABRIL","MAIO","JUNHO","JULHO"]

def hr_of(ws):
    for r in range(1,12):
        for c in range(1,ws.max_column+1):
            if norm(ws.cell(r,c).value).upper().startswith("PROCEDIMENTO"):
                return r
def maps(ws,hr):
    mc={};tc={}
    for c in range(1,ws.max_column+1):
        h=norm(ws.cell(hr,c).value).upper()
        if h in MONTHS: mc[h]=c
        elif "/TAB" in h:
            a=h.split("/")[0]
            for f in MONTHS:
                if f==a or f.startswith(a) or a.startswith(f[:3]): tc[f]=c;break
    return mc,tc
def rows(ws,hr):
    d={}
    for r in range(hr+1,ws.max_row+1):
        l=norm(ws.cell(r,1).value)
        if l=="":continue
        if l.upper().startswith("TOTAL PROCEDIMENTOS") or l.upper().startswith("TOTAL DE PROCEDIMENTOS"):break
        d[l.split("-")[0].strip()]=r
    return d

ANT=openpyxl.load_workbook("data/producao2026_v2.xlsx",data_only=True)
UPD=openpyxl.load_workbook("data/atualizada.xlsx",data_only=True)
MRG=openpyxl.load_workbook("data/producao_final_merged.xlsx",data_only=True)

CAPS=["CAPSIJ","CAPS II","CAPS AD III"]
all_ok=True

for name in MRG.sheetnames:
    mws=MRG[name]; mhr=hr_of(mws)
    mc,tc=maps(mws,mhr)
    heads=[norm(mws.cell(mhr,c).value) for c in range(1,mws.max_column+1) if norm(mws.cell(mhr,c).value)]
    print(f"\n### {name}")
    print("   headers:", heads)
    if name in CAPS:
        # 1) tab headers standardized
        tabheads=[h for h in heads if "/TAB" in h.upper()]
        expected=[f"{m}/TAB" for m in TAB_MONTHS]
        bad=[h for h in tabheads if h not in expected]
        print("   TAB headers:", tabheads)
        if bad or len(tabheads)!=7:
            print("   [!] PROBLEMA nos headers TAB:", bad); all_ok=False
        else:
            print("   [ok] 7 colunas TAB padronizadas por extenso")

        # 2) realizado unchanged vs updated
        uws=UPD[name]; uhr=hr_of(uws); umc,_=maps(uws,uhr)
        urows=rows(uws,uhr); mrows=rows(mws,mhr)
        changed=0
        for code,ur in urows.items():
            mr=mrows.get(code)
            if mr is None: print("   [!] código sumiu:",code);all_ok=False;continue
            for m in MONTHS:
                uc=umc.get(m); xc=mc.get(m)
                if uc is None or xc is None:continue
                uv=num(uws.cell(ur,uc).value); xv=num(mws.cell(mr,xc).value)
                if (uv or 0)!=(xv or 0):
                    changed+=1
                    print(f"   [!] REALIZADO alterado! {code} {m}: {uv}->{xv}")
        if changed==0: print("   [ok] REALIZADO idêntico à atualizada (nada foi alterado)")
        else: all_ok=False

        # 3) tab values match antiga
        aws=ANT[name]; ahr=hr_of(aws); _,atc=maps(aws,ahr); arows=rows(aws,ahr)
        mismatch=0; copied=0
        for code,mr in mrows.items():
            ar=arows.get(code)
            for m in TAB_MONTHS:
                xc=tc.get(m)
                if xc is None:continue
                xv=num(mws.cell(mr,xc).value)
                av=num(aws.cell(ar,atc[m]).value) if (ar and m in atc) else None
                if (av or 0)!=(xv or 0):
                    mismatch+=1
                    if mismatch<=5: print(f"   [!] TAB difere antiga {code} {m}: antiga={av} merged={xv}")
                elif xv is not None and xv!=0:
                    copied+=1
        if mismatch==0: print(f"   [ok] TAB idêntico à antiga ({copied} valores TAB não-zero conferidos)")
        else: all_ok=False

        # 4) confirm no TAB for AGOSTO/SETEMBRO
        if "AGOSTO" in tc or "SETEMBRO" in tc:
            print("   [!] inesperado: há coluna TAB para AGO/SET"); all_ok=False
        else:
            print("   [ok] AGOSTO e SETEMBRO sem coluna TAB (conforme esperado)")
    else:
        if any("/TAB" in h.upper() for h in heads):
            print("   [!] UA não deveria ter TAB"); all_ok=False
        else:
            print("   [ok] sem TAB (UA inalterada)")

print("\n"+("="*50))
print("RESULTADO FINAL:", "TUDO OK" if all_ok else "HÁ PROBLEMAS - REVISAR")

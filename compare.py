import openpyxl

def norm(s):
    return "" if s is None else str(s).strip()

def load(path):
    return openpyxl.load_workbook(path, data_only=True)

ANT = load("data/producao2026_v2.xlsx")   # a que voce ajeitou (tem TAB)
NEW = load("data/atualizada.xlsx")         # a atualizada (sem TAB, dados novos)

print("ANTIGA (v2) abas:", ANT.sheetnames)
print("NOVA (atualizada) abas:", NEW.sheetnames)
print()

def dump_headers(wb, tag):
    print(f"\n===== {tag} =====")
    for name in wb.sheetnames:
        ws = wb[name]
        # header row
        hr = None
        for r in range(1,12):
            for c in range(1, ws.max_column+1):
                if norm(ws.cell(r,c).value).upper().startswith("PROCEDIMENTO"):
                    hr = r; break
            if hr: break
        heads = [norm(ws.cell(hr,c).value) for c in range(1, ws.max_column+1) if norm(ws.cell(hr,c).value)!=""]
        print(f"  [{name}] header R{hr}: {heads}")

dump_headers(ANT, "ANTIGA v2 (com TAB)")
dump_headers(NEW, "NOVA atualizada")
